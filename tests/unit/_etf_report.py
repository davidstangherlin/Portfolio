"""Builds spreadsheets shaped like the ASX Investment Products report, for
the ETF tests (src/etf/asx_report.py). Two layouts, since the real one
isn't a published format: a title, a merged "Performance" group heading
over the period columns, fees as Excel percents and sizes in $m (layout
"grouped"); and a single heading row with whole-number percents and sizes
in dollars (layout "flat")."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

ETFS = [
    # code, name, issuer, type, exposure, benchmark, fee %, FUM $m, flows $m, spread %, 1m, 1y, 3y, 5y, 10y
    ("VAS", "Vanguard Australian Shares Index ETF", "Vanguard", "ETF", "Australian Equities", "S&P/ASX 300",
     0.07, 18500.5, 210.2, 0.02, 1.2, 12.5, 9.1, 8.4, 8.9),
    ("IOZ", "iShares Core S&P/ASX 200 ETF", "BlackRock", "ETF", "Australian Equities", "S&P/ASX 200",
     0.05, 6100.0, 55.0, 0.03, 1.1, 12.1, 8.9, 8.2, None),
    ("NDQ", "Betashares Nasdaq 100 ETF", "Betashares", "ETF", "Global Equities", "Nasdaq-100",
     0.48, 5200.0, -12.5, 0.05, -2.4, 24.3, 15.2, 17.8, None),
    ("HACK", "Betashares Global Cybersecurity ETF", "Betashares", "Active ETF", "Global Equities", None,
     0.67, 900.0, 3.0, 0.09, 3.3, 30.1, None, None, None),
    ("GOLD", "Global X Physical Gold", "Global X", "Structured Product", "Commodities", "Gold spot",
     0.40, 2500.0, 40.0, 0.06, 5.0, 35.0, 18.0, 12.0, 9.0),
]
LICS = [("AFI", "Australian Foundation Investment Co", "AFIC", "LIC", "Australian Equities")]


def build_report(folder: Path, name: str = "asx-investment-products-aug-2026.xlsx", layout: str = "grouped",
                 etfs=ETFS) -> Path:
    wb = Workbook()
    ws = wb.active
    if layout == "grouped":
        ws.title = "ETP Spotlight"
        ws.append(["ASX Investment Products - August 2026"])
        ws.append([])
        ws.append(["Fund details"] + [None] * 6 + ["Size and flows", None, None, "Performance (%)", None, None, None, None])
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
        ws.merge_cells(start_row=3, start_column=8, end_row=3, end_column=10)
        ws.merge_cells(start_row=3, start_column=11, end_row=3, end_column=15)
        ws.append(["ASX Code", "Fund Name", "Issuer", "Type", "Exposure", "Benchmark", "Management Fee",
                   "FUM ($m)", "Net Flows ($m)", "Avg Bid/Ask Spread", "1 Month", "1 Year", "3 Year", "5 Year", "10 Year"])
        for e in etfs:
            ws.append(list(e[:6]) + [e[6] / 100, e[7], e[8], e[9] / 100] + list(e[10:]))
            row = ws.max_row
            for col in (7, 10):
                ws.cell(row=row, column=col).number_format = "0.00%"
        for code, name_, issuer, type_, exposure in LICS:
            ws.append([code, name_, issuer, type_, exposure])
        ws.append(["Total", None, None, None, None, None, None, sum(e[7] for e in etfs)])
        lic = wb.create_sheet("LIC Spotlight")
        lic.append(["ASX Code", "Name", "Manager", "Type"])
        lic.append(["ARG", "Argo Investments", "Argo", "LIC"])
    else:
        ws.title = "Exchange Traded Products"
        ws.append(["ASX code", "ETP name", "Product issuer", "Product type", "Asset class", "Index",
                   "MER (%)", "Funds under management", "1 Mth Return (%)", "1 Yr Return (%)", "5 Yr Return (%)",
                   "Since Inception (%)", "Listing date", "Distribution frequency"])
        for e in etfs:
            ws.append(list(e[:6]) + [e[6], e[7] * 1_000_000, e[10], e[11], e[13], 7.5, "2009-05-04", "Quarterly"])
    path = Path(folder) / name
    wb.save(path)
    return path
