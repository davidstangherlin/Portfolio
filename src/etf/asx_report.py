"""The ASX Investment Products monthly report (docs/AS_BUILT.md §25): find
it, download it, read the ETF list out of it and load it.

The ASX publishes the report each month, by about the seventh business day
of the next month, as a PDF and a spreadsheet. The spreadsheet lists every
exchange traded product (ETFs, active ETFs, structured products) with its
issuer, category, fees, size, flows, spread and performance.

The spreadsheet's layout isn't a published format and has changed over
the years, so nothing here depends on fixed rows or column letters:
- the ETP sheets are picked by name (LIC, LIT, mFund, A-REIT and
  infrastructure sheets are skipped);
- the heading row is the one with an "ASX code" column, and a group
  heading above it (e.g. "Performance" over "1 Month", "1 Year") is
  carried into each column's label;
- each field is matched by the words in its label;
- percents and dollar units are read from the cell format and the
  label, with a check for fractions written as numbers.
Every column of every row is also kept as written (`raw`), so nothing is
lost if a label isn't recognised. `python -m src.etf.run_etfs --inspect
FILE` shows how a file is read without loading it.
"""

from __future__ import annotations

import logging
import re
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urljoin

from sqlalchemy import func, select, update
from sqlalchemy.dialects.postgresql import insert

from src.ingestion.dividend_history import add_months
from src.models import Company, EtfMonthly

logger = logging.getLogger(__name__)

REPORT_PAGE = "https://www.asx.com.au/issuers/investment-products/asx-investment-products-monthly-report"
# Where the PDFs are known to live; the spreadsheets are expected alongside.
# Only used if the report page itself can't be read.
REPORT_BASE = "https://www.asx.com.au/content/dam/asx/issuers/asx-investment-products-reports"
REPORT_DIR = Path(__file__).resolve().parents[2] / "data" / "asx_reports"
LATE_DAY = 15  # still missing after this day of the month: worth a warning

_MONTHS = ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec")
_MONTH_WORD = re.compile(r"(?<![a-z])(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[\s_\-]*'?(20\d\d|\d\d)(?!\d)", re.I)
_YEAR_MONTH = re.compile(r"(20\d\d)[\-_](0[1-9]|1[0-2])(?!\d)")
_CODE = re.compile(r"^[A-Z0-9]{3,6}$")

_SKIP_SHEETS = ("lic", "lit", "listed investment", "mfund", "m-fund", "reit", "infra", "unlisted", "issuer")
_ETP_SHEETS = ("etp", "etf", "exchange traded")
_SKIP_TYPES = ("lic", "lit", "listed investment", "mfund", "m-fund")
_SKIP_TYPE_EXACT = {"index"}   # benchmark index rows listed among the funds, not funds
# The report's own type abbreviations, spelt out.
TYPE_NAMES = {"SP": "Structured product", "ETF": "ETF", "Active": "Active ETF", "Complex": "Complex ETF"}
FUND_OF_FUNDS_MARK = "^"       # "an ETF that invests in whole or in part into another ETF admitted to ASX"
FUND_OF_FUNDS = "Invests in other ETFs"
_CODE_LABELS = {"asx code", "code", "ticker", "asx ticker", "etp code", "asx code ticker", "asx ticker code"}

PERCENT_FIELDS = ("mer_percent", "avg_spread_percent", "distribution_yield")
RETURN_FIELDS = ("return_1m", "return_3m", "return_6m", "return_1y", "return_3y", "return_5y", "return_10y",
                 "return_since_inception")
MONEY_FIELDS = ("fum_aud", "net_flows_aud", "value_traded_aud")
TEXT_FIELDS = ("fund_name", "issuer", "product_type", "category", "sub_category", "benchmark", "distribution_frequency")
# Not percent-formatted and the typical (median) value below this: the
# column holds fractions (0.0007 for 0.07%), so it's multiplied by 100.
# Judged on the median, so one fund with a 200% year can't hide it, and
# all the return columns are judged together, so a flat month can't either.
_FRACTION_MEDIAN = {"mer_percent": Decimal("0.05"), "avg_spread_percent": Decimal("0.02"),
                    "distribution_yield": Decimal("0.3"), "returns": Decimal("0.5")}
_PERIODS = {(1, "m"): "return_1m", (3, "m"): "return_3m", (6, "m"): "return_6m", (12, "m"): "return_1y",
            (1, "y"): "return_1y", (3, "y"): "return_3y", (5, "y"): "return_5y", (10, "y"): "return_10y"}
_PERIOD = re.compile(r"(\d+)\s*(m|mo|mth|mths|month|months|y|yr|yrs|year|years)(?![a-z])")
_PERIOD_FILLER = {"%", "return", "returns", "total", "tr", "pa", "p", "a", "annualised", "annualized", "to", "date"}
_NOT_RETURN = ("flow", "yield", "distribution", "spread", "volume", "value", "fee", "mer", "turnover", "trade", "cap")


class ReportError(Exception):
    """The file couldn't be read as an ASX Investment Products report."""


# ---------- finding and downloading the report ----------

def month_from_name(name: str) -> date | None:
    """The month a report covers, from its file name or link."""
    text = name.lower()
    m = _MONTH_WORD.search(text)
    if m:
        year = int(m.group(2))
        return date(year + 2000 if year < 100 else year, _MONTHS.index(m.group(1)[:3].lower()) + 1, 1)
    m = _YEAR_MONTH.search(text)
    return date(int(m.group(1)), int(m.group(2)), 1) if m else None


def report_links(html: str, base: str = REPORT_PAGE) -> list[tuple[date, str]]:
    """Every spreadsheet linked from the report page with a month in its
    name, newest first."""
    found: dict[str, date] = {}
    for href in re.findall(r"""href\s*=\s*["']([^"']+?\.xlsx)(?:\?[^"']*)?["']""", html, re.I):
        url = urljoin(base, href)
        month = month_from_name(url.rsplit("/", 1)[-1])
        if month:
            found[url] = month
    return sorted(((m, u) for u, m in found.items()), key=lambda x: x[0], reverse=True)


def guessed_urls(month: date) -> list[str]:
    mon, year = _MONTHS[month.month - 1], month.year
    names = (f"asx-investment-products-{mon}-{year}.xlsx", f"asx-investment-products-{mon}-{year}-abs.xlsx")
    return [f"{REPORT_BASE}/{year}/{folder}/{name}" for folder in ("excel", "xlsx") for name in names]


def _get(url: str) -> bytes | None:
    """The file at `url`, or None. Uses curl_cffi (installed with yfinance)
    with a browser's fingerprint, since the ASX site turns away plain
    scripted requests; falls back to the standard library."""
    try:
        from curl_cffi import requests as http
        response = http.get(url, impersonate="chrome", timeout=60)
        return response.content if response.status_code == 200 else None
    except ImportError:
        from urllib.request import Request, urlopen
        try:
            with urlopen(Request(url, headers={"User-Agent": "Mozilla/5.0"}), timeout=60) as response:
                return response.read()
        except OSError:
            return None
    except Exception:
        logger.debug("Download failed: %s", url, exc_info=True)
        return None


def _is_xlsx(content: bytes | None) -> bool:
    return bool(content) and content[:2] == b"PK"  # a spreadsheet is a zip file


def saved_path(month: date) -> Path:
    return REPORT_DIR / f"asx-investment-products-{month:%Y-%m}.xlsx"


def download_newer(have: date | None, today: date, fetch=_get, folder: Path = REPORT_DIR) -> tuple[date, Path] | None:
    """Download the newest report if it's newer than `have`. Reads the
    report page for its links; if that fails, tries the expected
    addresses for the last two months. Returns (month, saved file)."""
    candidates: list[tuple[date, str]] = []
    page = fetch(REPORT_PAGE)
    if page:
        candidates = report_links(page.decode("utf-8", errors="replace"))
        if not candidates:
            logger.warning("The ASX report page had no spreadsheet links; trying the expected addresses")
    else:
        logger.warning("Couldn't read the ASX report page; trying the expected addresses")
    if not candidates:
        this_month = today.replace(day=1)
        for month in (add_months(this_month, -1), add_months(this_month, -2)):
            candidates.extend((month, url) for url in guessed_urls(month))
    for month, url in candidates:
        if have is not None and month <= have:
            break
        content = fetch(url)
        if _is_xlsx(content):
            path = folder / saved_path(month).name
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
            logger.info("Downloaded the ASX report for %s from %s", f"{month:%B %Y}", url)
            return month, path
    return None


def local_newer(have: date | None, folder: Path = REPORT_DIR) -> tuple[date, Path] | None:
    """The newest spreadsheet in the reports folder newer than `have`, so a
    file downloaded by hand is picked up the same way."""
    if not folder.is_dir():
        return None
    files = [(month_from_name(p.name), p) for p in folder.glob("*.xlsx")]
    files = sorted(((m, p) for m, p in files if m and (have is None or m > have)), reverse=True)
    return files[0] if files else None


# ---------- reading the spreadsheet ----------

@dataclass
class Column:
    index: int
    label: str          # as written, group heading first
    field: str | None   # what it was matched to, or None (kept in raw only)


@dataclass
class SheetRead:
    name: str
    header_row: int
    columns: list[Column]
    rows: int = 0
    notes: list[str] = field(default_factory=list)


@dataclass
class Report:
    month: date | None
    rows: dict[str, dict]          # code -> fields (+ "raw")
    sheets: list[SheetRead]
    skipped_sheets: list[str]


def _clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _norm(label: str) -> str:
    text = label.lower().replace("&", " and ")
    text = re.sub(r"[^a-z0-9%$ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def match_field(label: str) -> str | None:
    """Which etf_monthly column a heading feeds, by its words."""
    n = _norm(label)
    if not n:
        return None
    words = set(n.split())
    if n in _CODE_LABELS or (n.endswith(" code") and ("asx" in words or "etp" in words)):
        return "code"
    if "since inception" in n or n.endswith(" si") or n == "si" or "inception to date" in n:
        if not any(w in n for w in _NOT_RETURN):
            return "return_since_inception"
    period = _PERIOD.search(n)
    if period and not any(w in n for w in _NOT_RETURN):
        bare = [w for w in _PERIOD.sub(" ", n).split() if w not in _PERIOD_FILLER]
        if any(w in n for w in ("perf", "return")) or not bare:
            unit = "y" if period.group(2).startswith("y") else "m"
            key = _PERIODS.get((int(period.group(1)), unit))
            if key:
                return key
    if "frequency" in n:
        return "distribution_frequency"
    if "yield" in n:
        return "distribution_yield"
    if "spread" in n:
        return "avg_spread_percent"
    if "flow" in n:
        return "net_flows_aud"
    if "value traded" in n or "transacted value" in n or "turnover" in n or "trading value" in n or ("value" in words and "traded" in words):
        return "value_traded_aud"
    if "mer" in words or "icr" in words or "fee" in n or "management cost" in n or "expense" in n:
        return "mer_percent"
    if "fum" in words or "funds under management" in n or "market cap" in n or "net assets" in n \
            or ("fund" in words and "size" in words):
        return "fum_aud"
    if ("listing" in words or "listed" in words or "inception" in words or "list" in words) and "date" in words:
        return "listing_date"
    if "issuer" in words or "manager" in words or "provider" in words or "sponsor" in words:
        return "issuer"
    if "name" in words:
        return "fund_name"
    if "benchmark" in words or "index" in words:
        return "benchmark"
    if "sub" in words or "subcategory" in words or "sub category" in n:
        return "sub_category"
    if "exposure" in words or "asset class" in n or "category" in words or "classification" in words \
            or "sector" in words:
        return "category"
    if "type" in words or "structure" in words:
        return "product_type"
    return None


def _labels(block: list[list]) -> list[tuple[str, str]]:
    """(full label, own label) per column from the heading rows: a group
    heading (a merged cell, so written once at its left) runs right until
    the next, and is put in front of each column's own heading."""
    width = max(len(r) for r in block)
    filled = []
    for i, row in enumerate(block):
        lowest, last, out = i == len(block) - 1, "", []
        for c in range(width):
            text = _clean(row[c]) if c < len(row) else ""
            if text:
                last = text
            elif not lowest:
                text = last
            out.append(text)
        filled.append(out)
    return [(" ".join(dict.fromkeys(f[c] for f in filled if f[c])), filled[-1][c]) for c in range(width)]


def column_field(full: str, own: str) -> str | None:
    """A column's own heading decides, unless it's only a period ("1
    Year") or says nothing recognisable: then the group heading in front
    decides whether it's performance, flows or something else."""
    f = match_field(own) if own else None
    if f is None or f in RETURN_FIELDS:
        return match_field(full) if full else None
    return f


def _is_code(value) -> bool:
    """ASX codes are written in capitals; "Total" or "Average" rows aren't codes."""
    text = _clean(value).removesuffix(".AX")
    return bool(_CODE.match(text)) and text not in ("TOTAL", "AVERAGE", "MEDIAN")


def _find_header(rows: list[list], limit: int = 40) -> int | None:
    for i, row in enumerate(rows[:limit]):
        texts = [_norm(_clean(v)) for v in row if _clean(v)]
        if len(texts) >= 3 and any(t in _CODE_LABELS or (t.endswith(" code") and "asx" in t) for t in texts):
            return i
    return None


def _number(value, percent_format: bool = False) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(round(value, 10) if isinstance(value, float) else value))
        except InvalidOperation:
            return None
        if not number.is_finite():
            return None
        return number * 100 if percent_format else number
    text = _clean(value).replace(",", "").replace("$", "").replace("%", "").strip()
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return -number if negative else number


def _money_scale(label: str) -> Decimal:
    n = label.lower()
    if re.search(r"\$\s*b\b|\bbn\b|billion|\(b\)", n):
        return Decimal("1e9")
    if re.search(r"\$\s*m\b|\bmn\b|million|\(m\)|\$m", n):
        return Decimal("1e6")
    if re.search(r"\$\s*k\b|\('?000s?\)|thousand", n):
        return Decimal("1e3")
    return Decimal("1")


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d %b %Y", "%d-%b-%y", "%b-%y", "%b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _raw_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, Decimal):
        return float(value)
    return value


def _skip_type(value) -> bool:
    """A product type that isn't an ETP (an LIC, LIT or mFund listed on the same sheet)."""
    text = _norm(_clean(value))
    return text in _SKIP_TYPE_EXACT or any(re.search(rf"(?<![a-z]){re.escape(s)}(?![a-z])", text) for s in _SKIP_TYPES)


def _sheet_wanted(name: str, all_names: list[str]) -> bool:
    n = name.lower()
    if any(s in n for s in _SKIP_SHEETS):
        return False
    named_etp = [x for x in all_names if any(s in x.lower() for s in _ETP_SHEETS)
                 and not any(s in x.lower() for s in _SKIP_SHEETS)]
    return not named_etp or name in named_etp


def _read_sheet(name: str, cells: list[list], formats: list[list]) -> tuple[SheetRead, dict[str, dict]] | None:
    header = _find_header(cells)
    if header is None:
        return None
    block_start = header
    if header > 0 and sum(1 for v in cells[header - 1] if _clean(v)) >= 2:
        block_start = header - 1  # a group heading row above
    first_data = header + 1
    code_col = next(i for i, v in enumerate(cells[header]) if match_field(_clean(v)) == "code")
    if first_data < len(cells) and not _is_code(cells[first_data][code_col] if code_col < len(cells[first_data]) else None) \
            and sum(1 for v in cells[first_data] if _clean(v)) >= 2:
        first_data += 1  # a sub-heading row below (e.g. "1 Month", "1 Year")
    labels = _labels(cells[block_start:first_data])

    columns, used = [], set()
    for i, (label, own) in enumerate(labels):
        f = "code" if i == code_col else column_field(label, own)
        if (f == "code" and i != code_col) or (f and f in used):
            f = None  # the first column for each field wins
        if f:
            used.add(f)
        columns.append(Column(i, label, f))
    read = SheetRead(name, header + 1, columns)

    data_rows = [(r, cells[r], formats[r]) for r in range(first_data, len(cells))]
    by_field = {c.field: c for c in columns if c.field}
    type_col = by_field.get("product_type")
    category_col = by_field.get("category")
    out: dict[str, dict] = {}
    section = None
    for _, row, fmt in data_rows:
        filled = [v for v in row if _clean(v)]
        if len(filled) == 1 and isinstance(filled[0], str) and not _is_code(filled[0]):
            section = _clean(filled[0])  # a section heading such as "Equity - Australia"
            continue
        if code_col >= len(row) or not _is_code(row[code_col]):
            continue
        if type_col and type_col.index < len(row) and _skip_type(row[type_col.index]):
            continue
        code = _clean(row[code_col]).removesuffix(".AX")
        item: dict = {"raw": {}}
        for col in columns:
            value = row[col.index] if col.index < len(row) else None
            is_pct = col.index < len(fmt) and "%" in (fmt[col.index] or "")
            if col.label and value is not None and _clean(value) != "":
                item["raw"][col.label] = _raw_value(value)
            if not col.field or col.field == "code":
                continue
            if col.field in TEXT_FIELDS:
                text = _clean(value)
                item[col.field] = text or None
            elif col.field == "listing_date":
                item[col.field] = _as_date(value)
            elif col.field in MONEY_FIELDS:
                number = _number(value)
                item[col.field] = number * _money_scale(col.label) if number is not None else None
            else:
                item[col.field] = _number(value, is_pct)
                item.setdefault("_pct_format", set())
                if is_pct:
                    item["_pct_format"].add(col.field)
        if FUND_OF_FUNDS_MARK in {_clean(v) for i, v in enumerate(row) if i < len(columns) and not columns[i].label}:
            item["raw"][FUND_OF_FUNDS] = True
        if category_col is None and section:
            item["category"] = section
        if item.get("product_type"):
            item["product_type"] = TYPE_NAMES.get(item["product_type"], item["product_type"])
        out[code] = item
    read.rows = len(out)
    if category_col is None and any(r.get("category") for r in out.values()):
        read.notes.append("category: taken from the section headings between the rows")
    _fix_fractions(read, out)
    _fix_money(read, out)
    for item in out.values():
        item.pop("_pct_format", None)
    return read, out


def _fix_fractions(read: SheetRead, rows: dict[str, dict]) -> None:
    for group, limit in _FRACTION_MEDIAN.items():
        fields = RETURN_FIELDS if group == "returns" else (group,)
        values = [abs(r[f]) for r in rows.values() for f in fields
                  if r.get(f) is not None and f not in r.get("_pct_format", ())]
        if len(values) >= 5 and statistics.median(values) < limit:
            for r in rows.values():
                for f in fields:
                    if r.get(f) is not None and f not in r.get("_pct_format", ()):
                        r[f] = r[f] * 100
            read.notes.append(f"{'returns' if group == 'returns' else group}: written as fractions, multiplied by 100")


def _fix_money(read: SheetRead, rows: dict[str, dict]) -> None:
    """Fund size with no unit in its label but a median under $100,000
    can only be in millions."""
    col = next((c for c in read.columns if c.field == "fum_aud"), None)
    values = [r["fum_aud"] for r in rows.values() if r.get("fum_aud")]
    if col and values and _money_scale(col.label) == 1 and statistics.median(values) < 100_000:
        for f in MONEY_FIELDS:
            for r in rows.values():
                if r.get(f) is not None:
                    r[f] = r[f] * Decimal("1e6")
        read.notes.append("fund size, flows and value traded: no unit given but sized in millions, multiplied by 1,000,000")


def _month_in_cells(sheets: list[list[list]]) -> date | None:
    for cells in sheets:
        for row in cells[:10]:
            for v in row:
                if isinstance(v, str) and re.search(r"(january|february|march|april|may|june|july|august|september|"
                                                    r"october|november|december)\s+20\d\d", v, re.I):
                    return month_from_name(v)
    return None


def read_report(path: Path, month: date | None = None) -> Report:
    """Read an ASX Investment Products spreadsheet into one entry per ETF
    code. Raises ReportError if no ETP sheet with an ASX code column is
    found."""
    try:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # not a spreadsheet, or corrupt
        raise ReportError(f"{Path(path).name} can't be opened as a spreadsheet: {exc}") from None
    names = list(book.sheetnames)
    sheets, skipped, rows, all_cells = [], [], {}, []
    for name in names:
        if not _sheet_wanted(name, names):
            skipped.append(name)
            continue
        ws = book[name]
        cells, formats = [], []
        for row in ws.iter_rows():
            cells.append([c.value for c in row])
            formats.append([getattr(c, "number_format", None) for c in row])
        all_cells.append(cells)
        result = _read_sheet(name, cells, formats)
        if result is None:
            skipped.append(name)
            continue
        read, found = result
        sheets.append(read)
        for code, item in found.items():
            if code in rows:  # listed on two sheets: keep what each adds
                for k, v in item.items():
                    if k == "raw":
                        rows[code]["raw"].update(v)
                    elif rows[code].get(k) is None:
                        rows[code][k] = v
            else:
                rows[code] = item
    book.close()
    if not rows:
        raise ReportError(f"No ETP list found in {Path(path).name}: no sheet has an ASX code column "
                          f"(sheets: {', '.join(names)})")
    return Report(month or month_from_name(Path(path).name) or _month_in_cells(all_cells), rows, sheets, skipped)


def describe(report: Report) -> str:
    """What --inspect prints: how each sheet was read."""
    lines = [f"Report month: {report.month:%B %Y}" if report.month else "Report month: not found (use --month)"]
    for s in report.sheets:
        lines.append(f"\nSheet '{s.name}': heading on row {s.header_row}, {s.rows} ETPs")
        for c in s.columns:
            if c.label:
                lines.append(f"  {c.label[:60]:<60} -> {c.field or '(kept in raw only)'}")
        lines.extend(f"  note: {n}" for n in s.notes)
    if report.skipped_sheets:
        lines.append(f"\nSkipped sheets: {', '.join(report.skipped_sheets)}")
    mapped = {c.field for s in report.sheets for c in s.columns if c.field} | \
        {k for item in report.rows.values() for k, v in item.items() if v is not None}
    missing = [f for f in ("fund_name", "issuer", "category", "mer_percent", "fum_aud", "return_1y") if f not in mapped]
    if missing:
        lines.append(f"\nNot found: {', '.join(missing)}")
    sample = list(report.rows.items())[:3]
    for code, item in sample:
        shown = {k: (str(v) if isinstance(v, Decimal) else v) for k, v in item.items() if k != "raw" and v is not None}
        lines.append(f"\n{code}: {shown}")
    return "\n".join(lines)


# ---------- loading ----------

@dataclass
class LoadResult:
    month: date
    etfs: int
    added: list[str]
    reclassified: list[str]
    deactivated: list[str]
    reactivated: list[str]


def _etf_company(session, code: str, name: str | None) -> tuple[Company, str | None]:
    company = session.execute(select(Company).where(Company.asx_code == code)).scalar_one_or_none()
    if company is None:
        company = Company(ticker=f"{code}.AX", asx_code=code, company_name=(name or code)[:255], security_type="ETF",
                          trading_currency="AUD", financial_currency="AUD", is_active=True)
        session.add(company)
        session.flush()
        return company, "added"
    change = None
    if company.security_type != "ETF":
        company.security_type, change = "ETF", "reclassified"
    elif not company.is_active:
        company.is_active, change = True, "reactivated"
    if name and company.company_name != name[:255]:
        company.company_name = name[:255]
    return company, change


def load_report(session, path: Path, month: date | None = None) -> LoadResult:
    """Load one report: every ETP becomes (or stays) an ETF in `companies`
    with that month's row in `etf_monthly`. When it's the newest report
    loaded, ETFs it no longer lists are marked inactive (their history is
    kept). Re-loading a month replaces that month's rows."""
    report = read_report(path, month)
    if report.month is None:
        raise ReportError(f"Can't tell which month {Path(path).name} covers; give it with --month YYYY-MM")
    newest = session.execute(select(func.max(EtfMonthly.report_month))).scalar_one()
    added, reclassified, reactivated = [], [], []
    for code, item in sorted(report.rows.items()):
        company, change = _etf_company(session, code, item.get("fund_name"))
        {"added": added, "reclassified": reclassified, "reactivated": reactivated}.get(change, []).append(code)
        values = {k: v for k, v in item.items() if k in EtfMonthly.__table__.columns}
        values.update(company_id=company.company_id, report_month=report.month, source_file=Path(path).name)
        stmt = insert(EtfMonthly).values(**values)
        session.execute(stmt.on_conflict_do_update(
            index_elements=[EtfMonthly.company_id, EtfMonthly.report_month],
            set_={k: stmt.excluded[k] for k in values if k not in ("company_id", "report_month")}
                | {"loaded_at": func.current_timestamp()}))
    deactivated = []
    if newest is None or report.month >= newest:
        gone = session.execute(select(Company).where(
            Company.security_type == "ETF", Company.is_active.is_(True),
            Company.asx_code.not_in(list(report.rows)))).scalars().all()
        deactivated = sorted(c.asx_code for c in gone)
        if gone:
            session.execute(update(Company).where(Company.company_id.in_([c.company_id for c in gone]))
                            .values(is_active=False))
    if reclassified:
        logger.warning("Now treated as ETFs (were shares): %s", ", ".join(reclassified))
    return LoadResult(report.month, len(report.rows), added, reclassified, deactivated, reactivated)


def latest_loaded(session) -> date | None:
    return session.execute(select(func.max(EtfMonthly.report_month))).scalar_one()


def expected_month(today: date) -> date:
    """The newest report that could be out: last month's."""
    return add_months(today.replace(day=1), -1)


def ensure_latest(session, today: date, fetch=_get, folder: Path = REPORT_DIR) -> LoadResult | None:
    """The nightly step: if last month's report isn't loaded yet, load it
    from the reports folder (a file saved by hand) or download it. Does
    nothing, without touching the network, once it's loaded."""
    have = latest_loaded(session)
    if have is not None and have >= expected_month(today):
        return None
    found = local_newer(have, folder) or download_newer(have, today, fetch, folder)
    if found is None:
        log = logger.warning if today.day > LATE_DAY else logger.info
        log("ASX report for %s not available yet (have %s). Download it from %s into %s if this persists.",
            f"{expected_month(today):%B %Y}", f"{have:%B %Y}" if have else "none", REPORT_PAGE, folder)
        return None
    month, path = found
    return load_report(session, path, month)
